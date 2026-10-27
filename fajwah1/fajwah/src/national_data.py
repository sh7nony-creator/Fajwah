"""National survey tables only; no regional computations or imputed values."""
from dataclasses import dataclass
from math import isfinite
import pandas as pd
from openpyxl import load_workbook
from .data import RAW
from .validation import validate_checksums

FILE = RAW / 'elderly_by_region_2025.xlsx.xlsx'
SOURCE = 'الهيئة العامة للإحصاء · مسح كبار السن 2025، كما ورد في ملف المشروع'
NATIONALITIES = {'all': 'جميع الجنسيات', 'saudi': 'سعوديون', 'non_saudi': 'غير سعوديين'}
SEXES = {'all': 'الجميع', 'male': 'ذكور', 'female': 'إناث'}

@dataclass(frozen=True)
class Indicator:
    label: str
    sheet: str
    table_number: int
    first_row: int
    last_row: int
    total_row: int | None
    unit: str
    keyword: str
    description: str
    caveat: str
    saudi_only: bool = False
    year: int = 2025

CATALOG = {
    'income': Indicator('مصادر الدخل', 'المصدر_9', 9, 5, 9, 10, 'نسبة مئوية', 'مصدر الأجور',
        'التوزيع النسبي حسب مصدر الأجور الوارد في الجدول؛ لا يعرض مبالغ الدخل أو متوسطه.',
        'العنوان الأصلي هو «مصدر الأجور». لا يمثل هذا الجدول قياسًا شاملًا لكل مصادر دخل الأسرة.'),
    'housing': Indicator('نوع السكن وحيازته', 'المصدر_4', 4, 5, 10, 11, 'فرد', 'حيازة المسكن',
        'أعداد كبار السن السعوديين حسب حيازة المسكن والجنس، وليست أعداد مساكن.',
        'المتاح حيازة المسكن فقط، وليس نوع المبنى (شقة/فيلا). الأعداد وطنية وللسعوديين فقط؛ لا تُحوّل إلى نسب في هذه الصفحة.', True),
    'financial': Indicator('التعثر المالي', 'المصدر_31', 32, 5, 7, 8, 'نسبة مئوية', 'التعثر المالي',
        'التوزيع النسبي لدرجات الصعوبة المالية كما ورد في المسح.',
        'درجات وصفية في المسح، وليست بيانات تعثر ائتماني أو مبالغ ديون. رقم الجدول الداخلي 32 رغم اسم الورقة المصدر_31.'),
    'chronic': Indicator('الأمراض المزمنة', 'المصدر_11', 11, 5, 7, 8, 'نسبة مئوية', 'الامراض المزمنة',
        'التوزيع النسبي لحالة المعاناة: لا مرض، مرض مزمن واحد، أو عدة أمراض مزمنة.',
        'استُخدم جدول الحالة العامة. لم يُستخدم جدول أنواع الأمراض (المصدر_12) لأن مقامه يحتاج توضيحًا؛ ولا تجمع أنواع الأمراض المتداخلة.'),
    'health': Indicator('الصحة', 'المصدر_10', 10, 5, 9, 10, 'نسبة مئوية', 'التقييم الذاتي',
        'التقييم الذاتي للوضع الصحي، من ممتاز إلى سيئ جدًا.',
        'يعكس تقييم المشاركين الذاتي، وليس تشخيصًا طبيًا أو تفسيرًا لأسباب الحالة الصحية.'),
    'digital': Indicator('الاندماج الرقمي', 'المصدر_32', 33, 5, 7, None, 'نسبة مئوية', 'الاندماج الرقمي',
        'نسب الوصول الرقمي والاستخدام الرقمي والاستفادة الرقمية كما نُشرت في الجدول.',
        'هذه أبعاد منفصلة وليست فئات توزيع؛ لا تُجمع ولا يُحسب متوسطها. مقامات الأبعاد وتعريفاتها التشغيلية التفصيلية غير مرفقة، فلا تُفسر الفروق بينها كمسار تحويل. رقم الجدول الداخلي 33.'),
    'life': Indicator('الرضا عن الحياة', 'المصدر_29', 29, 5, 9, 10, 'نسبة مئوية', 'الرضا العام عن الحياة',
        'التوزيع النسبي لمستوى الرضا العام عن الحياة، من راضٍ جدًا إلى غير راضٍ جدًا.',
        'إجابات ذاتية في المسح، ولا تثبت علاقة سببية مع الصحة أو الدخل أو المؤشرات الأخرى.'),
}

def parse_sheet(sheet, key):
    """Explicit row/column contracts preserve the published denominators."""
    spec = CATALOG[key]
    title_row = 2 if spec.saudi_only else 1
    title = str(sheet.cell(title_row, 1).value)
    if spec.keyword not in title:
        raise ValueError(f'تغير عنوان الجدول: {spec.sheet}')
    source_row = spec.last_row + (2 if spec.total_row else 1)
    source = str(sheet.cell(source_row, 1).value)
    if '2025' not in source or 'الهيئة العامة للإحصاء' not in source:
        raise ValueError(f'مصدر أو سنة غير متوقعة: {spec.sheet}')
    number_row = 3 if spec.saudi_only else 2
    if str(spec.table_number) not in str(sheet.cell(number_row, 1).value):
        raise ValueError(f'تغير رقم الجدول: {spec.sheet}')
    columns = []
    if spec.saudi_only:
        if [sheet.cell(4, c).value for c in range(2, 5)] != ['ذكر', 'أنثى', 'الإجمالي']:
            raise ValueError('رؤوس جدول السكن غير مطابقة')
        columns = [('saudi', sex, col) for sex, col in [('male', 2), ('female', 3), ('all', 4)]]
    else:
        if [sheet.cell(3, c).value for c in [2, 5, 8]] != ['سعودي', 'غير سعودي', 'الإجمالي العام']:
            raise ValueError(f'رؤوس الجنسية غير مطابقة: {spec.sheet}')
        for nationality, start in [('saudi', 2), ('non_saudi', 5), ('all', 8)]:
            if [sheet.cell(4, c).value for c in range(start, start+3)] != ['ذكر', 'أنثى', 'الإجمالي']:
                raise ValueError(f'رؤوس الجنس غير مطابقة: {spec.sheet}')
            columns.extend((nationality, sex, start+offset) for offset, sex in enumerate(['male', 'female', 'all']))
    def number(row, col):
        value = sheet.cell(row, col).value
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not isfinite(value):
            raise ValueError(f'قيمة ناقصة أو غير رقمية: {spec.sheet}!{sheet.cell(row,col).coordinate}؛ لم تُستبدل بصفر')
        if value < 0 or (spec.unit == 'نسبة مئوية' and value > 100):
            raise ValueError(f'قيمة خارج النطاق: {spec.sheet}!{sheet.cell(row,col).coordinate}')
        if spec.unit == 'فرد' and value != int(value):
            raise ValueError('عدد أفراد غير صحيح')
        return value
    records = []
    for row in range(spec.first_row, spec.last_row+1):
        category = sheet.cell(row, 1).value
        if not isinstance(category, str) or not category.strip():
            raise ValueError(f'فئة مفقودة: {spec.sheet}!A{row}')
        for nationality, sex, col in columns:
            records.append(dict(category=category, nationality=nationality, sex=sex,
                value=number(row, col), excel_cell=sheet.cell(row,col).coordinate, sheet=spec.sheet,
                year=spec.year, unit=spec.unit))
    frame = pd.DataFrame(records)
    if frame.duplicated(['category','nationality','sex']).any():
        raise ValueError('فئات مكررة في الجدول الوطني')
    if spec.total_row:
        if sheet.cell(spec.total_row, 1).value != 'الإجمالي':
            raise ValueError(f'صف إجمالي مفقود: {spec.sheet}')
        for nationality, sex, col in columns:
            total = number(spec.total_row, col)
            values = frame.loc[(frame.nationality == nationality) & (frame.sex == sex), 'value']
            tolerance = len(values)*.05 + 1e-8 if spec.unit == 'نسبة مئوية' else 1e-8
            if abs(values.sum()-total) > tolerance or (spec.unit == 'نسبة مئوية' and abs(total-100) > 1e-8):
                raise ValueError(f'عدم تطابق مجموع الفئات والإجمالي: {spec.sheet}')
    if spec.saudi_only:
        pivot = frame.pivot(index='category', columns='sex', values='value')
        if not ((pivot['male'] + pivot['female']) == pivot['all']).all():
            raise ValueError('عدم تطابق إجمالي الجنسين في السكن')
    return frame

def load_national_table(key):
    validate_checksums()
    workbook = load_workbook(FILE, data_only=True, read_only=True)
    try:
        return parse_sheet(workbook[CATALOG[key].sheet], key)
    finally:
        workbook.close()

def select_values(frame, key, nationality='all', sex='all'):
    if nationality not in NATIONALITIES or sex not in SEXES:
        raise ValueError('اختيار فئة غير صالح')
    if CATALOG[key].saudi_only and nationality != 'saudi':
        raise ValueError('بيانات السكن للسعوديين فقط')
    selected = frame.loc[(frame.nationality == nationality) & (frame.sex == sex)].copy()
    if selected.empty or selected.value.isna().any():
        raise ValueError('لا توجد قيم صالحة للفئة المختارة')
    return selected.reset_index(drop=True)
