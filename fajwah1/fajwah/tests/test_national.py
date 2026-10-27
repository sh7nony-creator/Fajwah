import json
import unittest
from pathlib import Path
from openpyxl import load_workbook
import numpy as np
from streamlit.testing.v1 import AppTest
from src.national_data import (CATALOG, FILE, NATIONALITIES, SEXES, parse_sheet, load_national_table, select_values)
from src.national_views import national_figure, describe_national, national_table
from src.data import load_regional_bundle
from src.presentation import prepare_regional

ROOT = Path(__file__).resolve().parents[1]

class NationalTests(unittest.TestCase):
    def test_every_value_matches_original_excel_cell(self):
        workbook = load_workbook(FILE, data_only=True)
        for key, spec in CATALOG.items():
            frame = load_national_table(key)
            self.assertFalse(frame.value.isna().any())
            self.assertEqual(frame.category.nunique(), spec.last_row-spec.first_row+1)
            for row in frame.itertuples():
                self.assertEqual(row.value, workbook[spec.sheet][row.excel_cell].value)
                self.assertEqual(row.year, 2025)
                self.assertEqual(row.unit, spec.unit)
        workbook.close()

    def test_charts_tables_and_descriptions_all_filters(self):
        for key, spec in CATALOG.items():
            frame = load_national_table(key)
            nationalities = ['saudi'] if spec.saudi_only else list(NATIONALITIES)
            for nationality in nationalities:
                for sex in SEXES:
                    selected = select_values(frame,key,nationality,sex)
                    chart = national_figure(selected,key)
                    np.testing.assert_array_equal(chart.data[0].x, selected.value)
                    self.assertEqual(list(chart.data[0].y), selected.category.tolist())
                    np.testing.assert_array_equal(national_table(selected)['القيمة'], selected.value)
                    self.assertTrue(describe_national(selected,key))

    def test_housing_is_count_and_saudi_only(self):
        frame = load_national_table('housing')
        self.assertEqual(set(frame.nationality), {'saudi'})
        self.assertEqual(select_values(frame,'housing','saudi','all').value.sum(), 1213333)
        with self.assertRaises(ValueError):
            select_values(frame,'housing','all')
        self.assertEqual(CATALOG['housing'].unit,'فرد')

    def test_explicit_zero_preserved_and_missing_rejected(self):
        workbook = load_workbook(FILE,data_only=True)
        sheet = workbook[CATALOG['life'].sheet]
        frame = parse_sheet(sheet,'life')
        zero = frame.loc[frame.excel_cell == 'F9','value']
        self.assertEqual(zero.iloc[0], 0)
        sheet['F9'] = None
        with self.assertRaisesRegex(ValueError,'لم تُستبدل بصفر'):
            parse_sheet(sheet,'life')
        workbook.close()

    def test_published_totals_and_nonadditive_dimensions(self):
        for key,spec in CATALOG.items():
            frame = load_national_table(key)
            if spec.total_row and spec.unit == 'نسبة مئوية':
                for _, values in frame.groupby(['nationality','sex']):
                    self.assertAlmostEqual(values.value.sum(),100,delta=len(values)*.05+1e-8)
        self.assertIsNone(CATALOG['digital'].total_row)
        digital = select_values(load_national_table('digital'),'digital')
        self.assertGreater(digital.value.sum(),100) # Published dimensions must not be normalized.
        np.testing.assert_allclose(digital.value, [46.5,84.4,71.0])

    def test_bad_percentage_and_header_fail(self):
        workbook = load_workbook(FILE,data_only=True)
        sheet = workbook[CATALOG['income'].sheet]
        sheet['J5'] = 101
        with self.assertRaises(ValueError):
            parse_sheet(sheet,'income')
        sheet['J5'] = 16.8
        sheet['B3'] = 'غير معروف'
        with self.assertRaises(ValueError):
            parse_sheet(sheet,'income')
        workbook.close()

    def test_internal_table_numbers_not_assumed_from_sheet_names(self):
        self.assertEqual(CATALOG['financial'].table_number,32)
        self.assertEqual(CATALOG['digital'].table_number,33)

    def test_national_page_all_sections(self):
        app = AppTest.from_file(str(ROOT/'app.py')).run(timeout=30)
        app.switch_page('pages/national.py').run(timeout=30)
        for key, spec in CATALOG.items():
            app.selectbox(key='national_section').select(key).run(timeout=30)
            self.assertEqual(len(app.exception),0)
            self.assertEqual(len(app.error),0)
            self.assertEqual(app.title[0].value,'المؤشرات الوطنية')
            nationality = 'saudi' if spec.saudi_only else app.selectbox(key='national_nationality').value
            expected = select_values(load_national_table(key),key,nationality,app.selectbox(key='national_sex').value)
            chart = json.loads(app.get('plotly_chart')[0].proto.spec)
            np.testing.assert_array_equal(chart['data'][0]['x'],expected.value)
            np.testing.assert_array_equal(app.dataframe[0].value['القيمة'], expected.value)
        # National filters do not overwrite regional metric/region state.
        app.switch_page('pages/regional.py').run(timeout=30)
        self.assertEqual(len(app.exception),0)
        self.assertEqual(len(app.error),0)
        self.assertEqual(app.radio(key='metric').value,'population')

    def test_national_filter_switches(self):
        app = AppTest.from_file(str(ROOT/'app.py')).run(timeout=30)
        app.switch_page('pages/national.py').run(timeout=30)
        for nationality in NATIONALITIES:
            app.selectbox(key='national_nationality').select(nationality).run(timeout=30)
            for sex in SEXES:
                app.selectbox(key='national_sex').select(sex).run(timeout=30)
                expected = select_values(load_national_table('income'),'income',nationality,sex)
                np.testing.assert_array_equal(app.dataframe[0].value['القيمة'],expected.value)
                self.assertEqual(len(app.exception),0)

    def test_no_regional_result_changed(self):
        import pandas as pd
        before, _ = load_regional_bundle()
        prepared = prepare_regional(before)
        for key in CATALOG:
            load_national_table(key)
        after, _ = load_regional_bundle()
        pd.testing.assert_frame_equal(before, after)
        pd.testing.assert_frame_equal(prepared, prepare_regional(after))
